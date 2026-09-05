import json
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from FundBoard.models import (
    Account,
    CompoundInterestScenario,
    Loan,
    LoanSimulationScenario,
)
from FundBoard.services.compound_interest import (
    CompoundInterestAssumptions,
    calculate_compound_interest,
)
from FundBoard.services.loan_calculator import LoanAssumptions, calculate_loan

AJAX = {"HTTP_X_REQUESTED_WITH": "XMLHttpRequest"}


def loan_assumptions(**overrides):
    values = {
        "principal": Decimal("100000"),
        "annual_rate": Decimal("12"),
        "calculation_mode": "PAYMENT",
        "duration_months": 12,
        "target_payment": None,
        "insurance_mode": "FIXED",
        "insurance_value": Decimal("10"),
        "insurance_basis": "INITIAL",
        "initial_fees": Decimal("500"),
        "deferment_months": 0,
        "deferment_type": "NONE",
        "one_off_prepayment": Decimal("0"),
        "one_off_prepayment_month": None,
        "recurring_prepayment": Decimal("0"),
        "recurring_prepayment_start_month": None,
        "start_date": date(2026, 1, 31),
    }
    values.update(overrides)
    return LoanAssumptions(**values)


def compound_assumptions(**overrides):
    values = {
        "initial_capital": Decimal("1000"),
        "periodic_contribution": Decimal("100"),
        "contribution_frequency": "MONTHLY",
        "contribution_timing": "END",
        "duration_years": 1,
        "annual_return": Decimal("12"),
        "annual_fees": Decimal("0"),
        "annual_inflation": Decimal("0"),
        "tax_mode": "NONE",
        "tax_rate": Decimal("0"),
        "start_date": date(2026, 1, 31),
    }
    values.update(overrides)
    return CompoundInterestAssumptions(**values)


class LoanCalculatorTests(SimpleTestCase):
    def test_fixed_rate_schedule_is_cent_rounded_and_closes_at_zero(self):
        result = calculate_loan(loan_assumptions())

        self.assertEqual(result.regular_payment, Decimal("8884.88"))
        self.assertEqual(result.duration_months, 12)
        self.assertEqual(result.total_interest, Decimal("6618.53"))
        self.assertEqual(result.total_insurance, Decimal("120.00"))
        self.assertEqual(result.total_cost, Decimal("7238.53"))
        self.assertEqual(result.schedule[0].payment_date, date(2026, 2, 28))
        self.assertEqual(result.schedule[-1].closing_balance, Decimal("0.00"))
        self.assertEqual(
            sum(row.principal_paid + row.prepayment for row in result.schedule),
            Decimal("100000.00"),
        )

    def test_zero_rate_adjusts_the_final_rounding_remainder(self):
        result = calculate_loan(loan_assumptions(annual_rate=Decimal("0")))

        self.assertEqual(result.regular_payment, Decimal("8333.33"))
        self.assertEqual(result.schedule[-1].scheduled_payment, Decimal("8333.37"))
        self.assertEqual(result.total_interest, Decimal("0.00"))

    def test_target_payment_calculates_duration(self):
        result = calculate_loan(
            loan_assumptions(
                calculation_mode="DURATION",
                duration_months=None,
                target_payment=Decimal("9000"),
            )
        )

        self.assertEqual(result.regular_payment, Decimal("9000.00"))
        self.assertEqual(result.duration_months, 12)
        self.assertEqual(result.schedule[-1].scheduled_payment, Decimal("7539.98"))

    def test_total_deferment_capitalizes_interest(self):
        result = calculate_loan(
            loan_assumptions(deferment_months=2, deferment_type="TOTAL")
        )

        self.assertEqual(result.schedule[0].closing_balance, Decimal("101000.00"))
        self.assertEqual(result.schedule[1].closing_balance, Decimal("102010.00"))
        self.assertEqual(result.regular_payment, Decimal("10770.43"))
        self.assertEqual(result.duration_months, 12)

    def test_prepayment_keeps_payment_and_shortens_duration(self):
        baseline = calculate_loan(loan_assumptions())
        result = calculate_loan(
            loan_assumptions(
                one_off_prepayment=Decimal("10000"),
                one_off_prepayment_month=3,
            )
        )

        self.assertEqual(result.duration_months, 11)
        self.assertLess(result.total_interest, baseline.total_interest)
        self.assertEqual(result.schedule[2].prepayment, Decimal("10000.00"))

    def test_recurring_prepayment_and_percentage_insurance_are_supported(self):
        result = calculate_loan(
            loan_assumptions(
                insurance_mode="PERCENT",
                insurance_value=Decimal("0.36"),
                insurance_basis="INITIAL",
                recurring_prepayment=Decimal("500"),
                recurring_prepayment_start_month=2,
            )
        )

        self.assertEqual(result.schedule[0].insurance, Decimal("30.00"))
        self.assertEqual(result.schedule[1].prepayment, Decimal("500.00"))
        self.assertEqual(result.duration_months, 12)
        self.assertEqual(result.total_insurance, Decimal("360.00"))

    def test_target_payment_below_interest_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "ne couvre pas les intérêts"):
            calculate_loan(
                loan_assumptions(
                    calculation_mode="DURATION",
                    duration_months=None,
                    target_payment=Decimal("1000"),
                    annual_rate=Decimal("24"),
                )
            )


class CompoundInterestCalculatorTests(SimpleTestCase):
    def test_monthly_contributions_and_nominal_growth(self):
        result = calculate_compound_interest(compound_assumptions())

        self.assertEqual(result.total_contributed, Decimal("2200.00"))
        self.assertEqual(result.final_nominal_value, Decimal("2395.07"))
        self.assertEqual(result.net_gains, Decimal("195.07"))
        self.assertEqual(len(result.timeline), 12)
        self.assertEqual(result.timeline[0].closing_value, Decimal("1110.00"))

    def test_fees_inflation_and_final_tax_are_explicit(self):
        result = calculate_compound_interest(
            compound_assumptions(
                periodic_contribution=Decimal("0"),
                contribution_frequency="ANNUAL",
                annual_fees=Decimal("1.2"),
                annual_inflation=Decimal("2.4"),
                tax_mode="FINAL_GAINS",
                tax_rate=Decimal("30"),
            )
        )

        self.assertEqual(result.gross_final_value, Decimal("1113.37"))
        self.assertEqual(result.estimated_fees, Decimal("12.74"))
        self.assertEqual(result.estimated_tax, Decimal("34.01"))
        self.assertEqual(result.final_nominal_value, Decimal("1079.36"))
        self.assertEqual(result.final_real_value, Decimal("1053.79"))
        self.assertEqual(result.timeline[-1].tax, Decimal("34.01"))

    def test_beginning_contributions_outgrow_end_contributions(self):
        beginning = calculate_compound_interest(
            compound_assumptions(contribution_timing="BEGIN")
        )
        end = calculate_compound_interest(compound_assumptions())

        self.assertGreater(beginning.final_nominal_value, end.final_nominal_value)

    def test_quarterly_calendar_has_four_end_of_period_contributions(self):
        result = calculate_compound_interest(
            compound_assumptions(contribution_frequency="QUARTERLY")
        )

        contribution_months = [
            row.month for row in result.timeline if row.contribution > 0
        ]
        self.assertEqual(contribution_months, [3, 6, 9, 12])
        self.assertEqual(result.total_contributed, Decimal("1400.00"))

    def test_empty_projection_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "capital initial ou un versement"):
            calculate_compound_interest(
                compound_assumptions(
                    initial_capital=Decimal("0"),
                    periodic_contribution=Decimal("0"),
                )
            )


class SimulationModelTests(TestCase):
    def test_scenarios_refuse_inconsistent_conditional_fields(self):
        user = get_user_model().objects.create_user(username="simulation-model")
        loan = LoanSimulationScenario(
            user=user,
            name="Durée manquante",
            principal=Decimal("100000"),
            annual_rate=Decimal("2"),
            calculation_mode=LoanSimulationScenario.CalculationMode.PAYMENT,
            duration_months=None,
        )
        investment = CompoundInterestScenario(
            user=user,
            name="Fiscalité incohérente",
            initial_capital=Decimal("1000"),
            periodic_contribution=Decimal("0"),
            tax_mode=CompoundInterestScenario.TaxMode.NONE,
            tax_rate=Decimal("30"),
        )

        with self.assertRaises(ValidationError):
            loan.full_clean()
        with self.assertRaises(ValidationError):
            investment.full_clean()


class SimulationViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.alice = get_user_model().objects.create_user(username="sim-alice")
        cls.bob = get_user_model().objects.create_user(username="sim-bob")
        cls.alice_loan = LoanSimulationScenario.objects.create(
            user=cls.alice,
            name="Prêt Alice",
            principal=Decimal("100000"),
            annual_rate=Decimal("2"),
            duration_months=120,
            insurance_value=Decimal("15"),
            start_date=date(2026, 1, 1),
        )
        cls.alice_loan_two = LoanSimulationScenario.objects.create(
            user=cls.alice,
            name="Prêt Alice bis",
            principal=Decimal("100000"),
            annual_rate=Decimal("3"),
            duration_months=120,
            start_date=date(2026, 1, 1),
        )
        cls.bob_loan = LoanSimulationScenario.objects.create(
            user=cls.bob,
            name="Prêt Bob privé",
            principal=Decimal("90000"),
            annual_rate=Decimal("2"),
            duration_months=120,
        )
        cls.alice_compound = CompoundInterestScenario.objects.create(
            user=cls.alice,
            name="Projet Alice",
            initial_capital=Decimal("1000"),
            periodic_contribution=Decimal("100"),
            duration_years=2,
        )
        cls.alice_compound_two = CompoundInterestScenario.objects.create(
            user=cls.alice,
            name="Projet Alice bis",
            initial_capital=Decimal("2000"),
            periodic_contribution=Decimal("50"),
            duration_years=2,
        )
        cls.bob_compound = CompoundInterestScenario.objects.create(
            user=cls.bob,
            name="Projet Bob privé",
            initial_capital=Decimal("1000"),
            periodic_contribution=Decimal("100"),
        )

    def setUp(self):
        self.client.force_login(self.alice)

    def test_list_and_details_are_private_and_use_central_charts(self):
        list_response = self.client.get(reverse("fundboard:simulations"))
        detail_response = self.client.get(
            reverse("fundboard:loan_simulation_detail", args=[self.alice_loan.pk])
        )

        self.assertEqual(list_response.status_code, 200)
        self.assertContains(list_response, "Prêt Alice")
        self.assertNotContains(list_response, "Prêt Bob privé")
        self.assertContains(detail_response, "Tableau d'amortissement")
        self.assertContains(detail_response, "MizzacCharts.mount")
        self.assertEqual(
            self.client.get(
                reverse("fundboard:loan_simulation_detail", args=[self.bob_loan.pk])
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.get(
                reverse(
                    "fundboard:compound_simulation_detail",
                    args=[self.bob_compound.pk],
                )
            ).status_code,
            404,
        )

    def test_create_loan_scenario_does_not_create_real_financial_data(self):
        before = (Account.objects.count(), Loan.objects.count())
        response = self.client.post(
            reverse("fundboard:add_loan_simulation"),
            {
                "name": "Nouveau scénario",
                "currency": "EUR",
                "principal": "200000",
                "annual_rate": "3.2",
                "calculation_mode": "PAYMENT",
                "duration_months": "240",
                "target_payment": "",
                "insurance_mode": "FIXED",
                "insurance_value": "25",
                "insurance_basis": "INITIAL",
                "initial_fees": "1000",
                "deferment_months": "0",
                "deferment_type": "NONE",
                "one_off_prepayment": "0",
                "one_off_prepayment_month": "",
                "recurring_prepayment": "0",
                "recurring_prepayment_start_month": "",
                "start_date": "2026-09-05",
            },
            **AJAX,
        )

        scenario = LoanSimulationScenario.objects.get(name="Nouveau scénario")
        self.assertRedirects(
            response,
            reverse("fundboard:loan_simulation_detail", args=[scenario.pk]),
        )
        self.assertEqual(scenario.user, self.alice)
        self.assertEqual((Account.objects.count(), Loan.objects.count()), before)

    def test_create_and_archive_compound_scenario(self):
        response = self.client.post(
            reverse("fundboard:add_compound_simulation"),
            {
                "name": "Projet long terme",
                "currency": "EUR",
                "initial_capital": "5000",
                "periodic_contribution": "200",
                "contribution_frequency": "MONTHLY",
                "contribution_timing": "END",
                "duration_years": "15",
                "annual_return": "6",
                "annual_fees": "0.5",
                "annual_inflation": "2",
                "tax_mode": "FINAL_GAINS",
                "tax_rate": "30",
                "start_date": "2026-09-05",
            },
            **AJAX,
        )
        scenario = CompoundInterestScenario.objects.get(name="Projet long terme")

        self.assertRedirects(
            response,
            reverse("fundboard:compound_simulation_detail", args=[scenario.pk]),
        )
        archive_response = self.client.post(
            reverse("fundboard:archive_compound_simulation", args=[scenario.pk]),
            **AJAX,
        )
        scenario.refresh_from_db()
        self.assertRedirects(archive_response, reverse("fundboard:simulations"))
        self.assertTrue(scenario.archived)

    def test_foreign_scenario_cannot_be_archived(self):
        response = self.client.post(
            reverse(
                "fundboard:archive_compound_simulation",
                args=[self.bob_compound.pk],
            ),
            **AJAX,
        )

        self.assertEqual(response.status_code, 404)
        self.bob_compound.refresh_from_db()
        self.assertFalse(self.bob_compound.archived)

    def test_comparison_accepts_two_owned_scenarios_and_rejects_foreign_one(self):
        url = reverse("fundboard:loan_simulation_compare")
        response = self.client.get(
            url,
            {"ids": [self.alice_loan.pk, self.alice_loan_two.pk]},
        )
        rejected = self.client.get(
            url,
            {"ids": [self.alice_loan.pk, self.bob_loan.pk]},
        )

        self.assertContains(response, "Prêt Alice bis")
        self.assertContains(response, "loan-comparison-series")
        self.assertContains(rejected, "Sélectionnez entre deux et quatre")
        self.assertNotContains(rejected, "Prêt Bob privé")

    def test_json_and_excel_exports_are_owned_and_complete(self):
        json_response = self.client.get(
            reverse(
                "fundboard:compound_simulation_export",
                args=[self.alice_compound.pk, "json"],
            )
        )
        payload = json.loads(json_response.content)
        excel_response = self.client.get(
            reverse(
                "fundboard:loan_simulation_export",
                args=[self.alice_loan.pk, "xlsx"],
            )
        )
        workbook = load_workbook(BytesIO(excel_response.content), read_only=True)

        self.assertEqual(payload["simulation_type"], "compound_interest")
        self.assertEqual(len(payload["timeline"]), 24)
        self.assertIn("final_real_value", payload["result"])
        self.assertEqual(
            workbook.sheetnames,
            ["Hypothèses", "Résultats", "Amortissement"],
        )
        self.assertEqual(
            self.client.get(
                reverse(
                    "fundboard:loan_simulation_export",
                    args=[self.bob_loan.pk, "json"],
                )
            ).status_code,
            404,
        )

    def test_excel_export_neutralizes_a_formula_like_scenario_name(self):
        self.alice_loan.name = "=HYPERLINK(\"https://invalid.test\")"
        self.alice_loan.save(update_fields=["name"])

        response = self.client.get(
            reverse(
                "fundboard:loan_simulation_export",
                args=[self.alice_loan.pk, "xlsx"],
            )
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)

        self.assertTrue(workbook["Hypothèses"]["B2"].value.startswith("'="))

    def test_compound_comparison_escapes_names_in_chart_json(self):
        self.alice_compound.name = "</script><script>alert(1)</script>"
        self.alice_compound.save(update_fields=["name"])

        response = self.client.get(
            reverse("fundboard:compound_simulation_compare"),
            {"ids": [self.alice_compound.pk, self.alice_compound_two.pk]},
        )

        self.assertContains(response, r"\u003C/script\u003E", html=False)
        self.assertNotContains(response, "<script>alert(1)</script>", html=False)
