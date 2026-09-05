import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from FundBoard.exports.paper_trading import (
    virtual_portfolio_excel_bytes,
    virtual_portfolio_json_bytes,
)
from FundBoard.models import (
    Account,
    Instrument,
    Position,
    Price,
    Transaction,
    VirtualCashEvent,
    VirtualCorporateAction,
    VirtualOrder,
    VirtualPortfolio,
    VirtualPosition,
)
from FundBoard.services.paper_trading import (
    add_to_virtual_watchlist,
    apply_virtual_corporate_action,
    cancel_virtual_order,
    clone_virtual_portfolio,
    create_virtual_snapshot,
    initialize_virtual_portfolio,
    place_virtual_order,
    process_open_virtual_orders,
    reset_virtual_portfolio,
    searchable_virtual_instruments,
    virtual_portfolio_valuation,
)


class PaperTradingFixtureMixin:
    @classmethod
    def setUpTestData(cls):
        cls.alice = get_user_model().objects.create_user(
            username="paper-alice",
            password="test-password",
        )
        cls.bob = get_user_model().objects.create_user(username="paper-bob")
        cls.stock = Instrument.objects.create(
            name="Action Démo",
            instrument_type=Instrument.Type.STOCK,
            ticker="DEMO.PA",
            currency="EUR",
        )
        cls.crypto = Instrument.objects.create(
            name="Bitcoin Démo",
            instrument_type=Instrument.Type.CRYPTO,
            ticker="BTC",
            currency="USD",
        )
        cls.bob_instrument = Instrument.objects.create(
            owner=cls.bob,
            name="Privé Bob",
            instrument_type=Instrument.Type.STOCK,
            ticker="BOB",
            currency="EUR",
        )

    def portfolio(self, **overrides):
        values = {
            "user": self.alice,
            "name": "Laboratoire",
            "base_currency": "EUR",
            "initial_cash": Decimal("10000"),
            "cash_balance": Decimal("10000"),
        }
        values.update(overrides)
        portfolio = VirtualPortfolio.objects.create(**values)
        initialize_virtual_portfolio(portfolio)
        return portfolio

    def quote(
        self,
        instrument=None,
        *,
        close="100",
        currency="EUR",
        observed_at=None,
        source="test",
        market_state=Price.MarketState.REGULAR,
        delayed=False,
        quality=Price.Quality.FRESH,
    ):
        instant = observed_at or timezone.now() - timedelta(minutes=1)
        return Price.objects.create(
            instrument=instrument or self.stock,
            observed_at=instant,
            close_price=Decimal(close),
            currency=currency,
            source=source,
            collected_at=instant,
            market_state=market_state,
            is_delayed=delayed,
            quality=quality,
        )


class PaperTradingModelTests(PaperTradingFixtureMixin, TestCase):
    def test_cross_user_instrument_and_incoherent_limit_are_rejected(self):
        portfolio = self.portfolio()
        watch = portfolio.watchlist_entries.model(
            portfolio=portfolio,
            instrument=self.bob_instrument,
        )
        order = VirtualOrder(
            portfolio=portfolio,
            instrument=self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.LIMIT,
            quantity=Decimal("1"),
        )

        with self.assertRaises(ValidationError):
            watch.full_clean()
        with self.assertRaises(ValidationError):
            order.full_clean()

    def test_virtual_models_have_no_relation_to_real_accounts_or_transactions(self):
        forbidden_targets = {Account, Position, Transaction}

        for model in (
            VirtualPortfolio,
            VirtualPosition,
            VirtualOrder,
            VirtualCashEvent,
            VirtualCorporateAction,
        ):
            related_targets = {
                field.related_model
                for field in model._meta.get_fields()
                if getattr(field, "related_model", None) is not None
            }
            self.assertTrue(forbidden_targets.isdisjoint(related_targets))


class PaperTradingExecutionTests(PaperTradingFixtureMixin, TestCase):
    def test_market_buy_uses_past_local_price_fees_and_never_touches_real_ledger(self):
        portfolio = self.portfolio(
            proportional_fee_rate=Decimal("0.1"),
            fixed_fee=Decimal("1"),
            spread_bps=Decimal("20"),
            slippage_bps=Decimal("5"),
        )
        add_to_virtual_watchlist(portfolio, self.stock, allow_fractional=False)
        now = timezone.now()
        past = self.quote(observed_at=now - timedelta(minutes=1), delayed=True)
        self.quote(
            close="999",
            observed_at=now + timedelta(minutes=1),
            source="future",
        )

        with patch("requests.sessions.Session.request") as network_call:
            order = place_virtual_order(
                portfolio,
                self.stock,
                side=VirtualOrder.Side.BUY,
                order_type=VirtualOrder.OrderType.MARKET,
                quantity=Decimal("10"),
                submitted_at=now,
            )

        order.refresh_from_db()
        portfolio.refresh_from_db()
        position = VirtualPosition.objects.get(portfolio=portfolio, instrument=self.stock)
        self.assertEqual(order.status, VirtualOrder.Status.EXECUTED)
        self.assertEqual(order.execution_unit_price, Decimal("100.150000000000"))
        self.assertEqual(order.gross_amount, Decimal("1001.50000000"))
        self.assertEqual(order.fees, Decimal("2.00000000"))
        self.assertEqual(order.price_observed_at, past.observed_at)
        self.assertTrue(order.price_is_delayed)
        self.assertEqual(portfolio.cash_balance, Decimal("8996.50000000"))
        self.assertEqual(position.average_unit_cost, Decimal("100.350000000000"))
        self.assertEqual(position.quantity, Decimal("10"))
        self.assertEqual(Account.objects.count(), 0)
        self.assertEqual(Position.objects.count(), 0)
        self.assertEqual(Transaction.objects.count(), 0)
        network_call.assert_not_called()

    def test_limit_order_waits_then_executes_on_a_new_timestamped_price(self):
        portfolio = self.portfolio()
        add_to_virtual_watchlist(portfolio, self.stock)
        first_time = timezone.now()
        self.quote(close="100", observed_at=first_time - timedelta(minutes=1))

        order = place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.LIMIT,
            quantity=Decimal("1"),
            limit_price=Decimal("90"),
            submitted_at=first_time,
        )
        self.assertEqual(order.status, VirtualOrder.Status.OPEN)
        self.assertIn("limite", order.status_message)

        second_time = first_time + timedelta(minutes=2)
        crossed = self.quote(
            close="89",
            observed_at=second_time - timedelta(seconds=1),
            source="later",
        )
        (result,) = process_open_virtual_orders(portfolio, as_of=second_time)

        self.assertEqual(result.status, VirtualOrder.Status.EXECUTED)
        self.assertEqual(result.price_observed_at, crossed.observed_at)
        self.assertEqual(result.execution_unit_price, Decimal("89.000000000000"))

    def test_closed_or_stale_equity_stays_open_and_can_be_cancelled(self):
        portfolio = self.portfolio()
        add_to_virtual_watchlist(portfolio, self.stock)
        now = timezone.now()
        self.quote(
            observed_at=now - timedelta(minutes=1),
            market_state=Price.MarketState.CLOSED,
        )
        order = place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("1"),
            submitted_at=now,
        )
        self.assertEqual(order.status, VirtualOrder.Status.OPEN)
        self.assertIn("séance régulière", order.status_message)

        cancelled = cancel_virtual_order(order)
        self.assertEqual(cancelled.status, VirtualOrder.Status.CANCELLED)
        with self.assertRaisesRegex(ValueError, "ordre ouvert"):
            cancel_virtual_order(cancelled)

        stale_portfolio = self.portfolio(name="Cours ancien")
        add_to_virtual_watchlist(stale_portfolio, self.stock)
        Price.objects.all().delete()
        self.quote(observed_at=now - timedelta(minutes=30))
        stale = place_virtual_order(
            stale_portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("1"),
            submitted_at=now,
        )
        self.assertEqual(stale.status, VirtualOrder.Status.OPEN)
        self.assertIn("trop ancien", stale.status_message)

    def test_crypto_is_24_7_and_fractional_by_default(self):
        portfolio = self.portfolio(
            name="Crypto",
            base_currency="USD",
            initial_cash=Decimal("1000"),
            cash_balance=Decimal("1000"),
        )
        entry = add_to_virtual_watchlist(portfolio, self.crypto)
        self.quote(
            self.crypto,
            close="200",
            currency="USD",
            market_state=Price.MarketState.CLOSED,
        )

        order = place_virtual_order(
            portfolio,
            self.crypto,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("0.5"),
        )

        self.assertTrue(entry.allow_fractional)
        self.assertEqual(order.status, VirtualOrder.Status.EXECUTED)

    def test_fractional_stock_and_overselling_are_rejected(self):
        portfolio = self.portfolio()
        add_to_virtual_watchlist(portfolio, self.stock, allow_fractional=False)
        self.quote()

        with self.assertRaisesRegex(ValueError, "quantité entière"):
            place_virtual_order(
                portfolio,
                self.stock,
                side=VirtualOrder.Side.BUY,
                order_type=VirtualOrder.OrderType.MARKET,
                quantity=Decimal("0.5"),
            )
        sell = place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.SELL,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("1"),
        )
        self.assertEqual(sell.status, VirtualOrder.Status.REJECTED)
        self.assertIn("Position virtuelle insuffisante", sell.status_message)

    def test_realized_gain_survives_a_full_position_close(self):
        portfolio = self.portfolio()
        add_to_virtual_watchlist(portfolio, self.stock)
        buy_time = timezone.now()
        self.quote(close="100", observed_at=buy_time - timedelta(seconds=1))
        place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("1"),
            submitted_at=buy_time,
        )
        sell_time = buy_time + timedelta(minutes=1)
        self.quote(
            close="120",
            observed_at=sell_time - timedelta(seconds=1),
            source="sell",
        )
        place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.SELL,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("1"),
            submitted_at=sell_time,
        )

        position = VirtualPosition.objects.get(portfolio=portfolio, instrument=self.stock)
        valuation = virtual_portfolio_valuation(portfolio, as_of=sell_time)
        self.assertEqual(position.quantity, Decimal("0"))
        self.assertEqual(position.realized_gain, Decimal("20.00000000"))
        self.assertEqual(valuation.realized_gain, Decimal("20.00"))
        self.assertEqual(valuation.absolute_gain, Decimal("20.00"))
        self.assertEqual(valuation.positions, ())


class PaperTradingLifecycleTests(PaperTradingFixtureMixin, TestCase):
    def held_portfolio(self):
        portfolio = self.portfolio()
        add_to_virtual_watchlist(portfolio, self.stock)
        now = timezone.now()
        self.quote(close="100", observed_at=now - timedelta(minutes=1))
        place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("10"),
            submitted_at=now,
        )
        return portfolio, now

    def test_dividend_and_split_update_only_the_virtual_journal(self):
        portfolio, now = self.held_portfolio()
        open_order = place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.LIMIT,
            quantity=Decimal("2"),
            limit_price=Decimal("50"),
            submitted_at=now,
        )
        dividend = VirtualCorporateAction.objects.create(
            portfolio=portfolio,
            instrument=self.stock,
            action_type=VirtualCorporateAction.Type.DIVIDEND,
            effective_at=now,
            dividend_per_unit=Decimal("2"),
        )
        split = VirtualCorporateAction.objects.create(
            portfolio=portfolio,
            instrument=self.stock,
            action_type=VirtualCorporateAction.Type.SPLIT,
            effective_at=now,
            split_ratio=Decimal("2"),
        )

        apply_virtual_corporate_action(dividend, as_of=now)
        apply_virtual_corporate_action(dividend, as_of=now)
        apply_virtual_corporate_action(split, as_of=now)

        portfolio.refresh_from_db()
        position = portfolio.virtual_positions.get(instrument=self.stock)
        open_order.refresh_from_db()
        self.assertEqual(portfolio.cash_balance, Decimal("9020.00000000"))
        self.assertEqual(position.dividend_income, Decimal("20.00000000"))
        self.assertEqual(position.quantity, Decimal("20"))
        self.assertEqual(position.average_unit_cost, Decimal("50.000000000000"))
        self.assertEqual(open_order.quantity, Decimal("4"))
        self.assertEqual(open_order.limit_price, Decimal("25.000000000000"))
        self.assertEqual(Account.objects.count(), 0)
        self.assertEqual(Transaction.objects.count(), 0)

    def test_future_corporate_action_remains_pending(self):
        portfolio, now = self.held_portfolio()
        action = VirtualCorporateAction.objects.create(
            portfolio=portfolio,
            instrument=self.stock,
            action_type=VirtualCorporateAction.Type.DIVIDEND,
            effective_at=now + timedelta(days=1),
            dividend_per_unit=Decimal("1"),
        )

        result = apply_virtual_corporate_action(action, as_of=now)

        self.assertEqual(result.status, VirtualCorporateAction.Status.PENDING)
        self.assertIn("futur", result.status_message)

    def test_benchmark_snapshot_clone_and_reset(self):
        benchmark = Instrument.objects.create(
            name="Indice Démo",
            instrument_type=Instrument.Type.INDEX,
            ticker="IDX",
            currency="EUR",
        )
        initial_time = timezone.now() - timedelta(minutes=2)
        self.quote(
            benchmark,
            close="100",
            observed_at=initial_time,
            source="benchmark-initial",
        )
        portfolio = self.portfolio(benchmark_instrument=benchmark)
        VirtualPortfolio.objects.filter(pk=portfolio.pk).update(
            created_at=initial_time + timedelta(minutes=1)
        )
        portfolio.refresh_from_db()
        current_time = timezone.now()
        self.quote(
            benchmark,
            close="110",
            observed_at=current_time - timedelta(seconds=1),
            source="benchmark-current",
        )
        add_to_virtual_watchlist(portfolio, self.stock)
        self.quote(close="100", observed_at=current_time - timedelta(seconds=1))
        place_virtual_order(
            portfolio,
            self.stock,
            side=VirtualOrder.Side.BUY,
            order_type=VirtualOrder.OrderType.MARKET,
            quantity=Decimal("10"),
            submitted_at=current_time,
        )

        snapshot = create_virtual_snapshot(portfolio, as_of=current_time)
        clone = clone_virtual_portfolio(portfolio, name="Copie indépendante")

        self.assertEqual(snapshot.benchmark_value, Decimal("11000.00000000"))
        self.assertEqual(clone.user, self.alice)
        self.assertEqual(clone.initial_cash, Decimal("10000.00000000"))
        self.assertEqual(clone.virtual_positions.get().quantity, Decimal("10"))
        self.assertEqual(clone.virtual_orders.count(), 0)
        self.assertEqual(clone.cash_events.get().event_type, VirtualCashEvent.Type.CLONE)

        reset_virtual_portfolio(portfolio)
        portfolio.refresh_from_db()
        self.assertEqual(portfolio.cash_balance, Decimal("10000.00000000"))
        self.assertEqual(portfolio.virtual_positions.count(), 0)
        self.assertEqual(portfolio.virtual_orders.count(), 0)
        self.assertEqual(portfolio.watchlist_entries.count(), 1)
        self.assertEqual(portfolio.cash_events.get().event_type, VirtualCashEvent.Type.RESET)

    def test_exports_are_versioned_complete_and_formula_safe(self):
        portfolio = self.portfolio(name="=HYPERLINK(1)")
        valuation = virtual_portfolio_valuation(portfolio)

        payload = json.loads(virtual_portfolio_json_bytes(portfolio, valuation))
        workbook = load_workbook(BytesIO(virtual_portfolio_excel_bytes(portfolio, valuation)))

        self.assertEqual(payload["schema_version"], "1.0")
        self.assertEqual(payload["resource"], "virtual_portfolio")
        self.assertIn("aucun appel courtier", payload["execution_policy"])
        self.assertEqual(
            set(workbook.sheetnames),
            {"Portefeuille", "Watchlist", "Positions", "Ordres", "Cash", "Événements", "Performance"},
        )
        self.assertEqual(workbook["Portefeuille"]["B2"].value, "'=HYPERLINK(1)")


class PaperTradingViewTests(PaperTradingFixtureMixin, TestCase):
    def setUp(self):
        self.client.force_login(self.alice)

    def test_create_list_detail_search_and_exports(self):
        create_url = reverse("fundboard:add_virtual_portfolio")
        response = self.client.post(
            create_url,
            {
                "name": "Bac à sable",
                "base_currency": "EUR",
                "initial_cash": "5000",
                "benchmark_instrument": "",
                "proportional_fee_rate": "0",
                "fixed_fee": "0",
                "spread_bps": "0",
                "slippage_bps": "0",
            },
        )
        portfolio = VirtualPortfolio.objects.get(user=self.alice, name="Bac à sable")

        self.assertRedirects(
            response,
            reverse("fundboard:virtual_portfolio_detail", args=[portfolio.pk]),
        )
        self.assertEqual(portfolio.cash_events.count(), 1)
        self.assertEqual(portfolio.performance_snapshots.count(), 1)
        list_response = self.client.get(reverse("fundboard:virtual_portfolios"))
        detail_response = self.client.get(
            reverse("fundboard:virtual_portfolio_detail", args=[portfolio.pk]),
            {"q": "Démo"},
        )
        self.assertContains(list_response, "PORTEFEUILLE VIRTUEL")
        self.assertContains(detail_response, "Action Démo")
        self.assertContains(detail_response, "aucun ordre réel")

        json_response = self.client.get(
            reverse(
                "fundboard:virtual_portfolio_export",
                args=[portfolio.pk, "json"],
            )
        )
        xlsx_response = self.client.get(
            reverse(
                "fundboard:virtual_portfolio_export",
                args=[portfolio.pk, "xlsx"],
            )
        )
        self.assertEqual(json_response.status_code, 200)
        self.assertEqual(json.loads(json_response.content)["resource"], "virtual_portfolio")
        self.assertEqual(xlsx_response.status_code, 200)
        self.assertTrue(xlsx_response.content.startswith(b"PK"))

    def test_user_cannot_view_mutate_or_export_another_portfolio(self):
        bob_portfolio = self.portfolio(user=self.bob, name="Secret Bob")

        protected_urls = (
            reverse("fundboard:virtual_portfolio_detail", args=[bob_portfolio.pk]),
            reverse("fundboard:edit_virtual_portfolio", args=[bob_portfolio.pk]),
            reverse("fundboard:reset_virtual_portfolio", args=[bob_portfolio.pk]),
            reverse(
                "fundboard:virtual_portfolio_export",
                args=[bob_portfolio.pk, "json"],
            ),
        )
        for url in protected_urls:
            response = self.client.get(
                url,
                HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            )
            self.assertEqual(response.status_code, 404)

    def test_watchlist_order_cancel_clone_and_reset_http_workflow(self):
        portfolio = self.portfolio()
        add_watch = self.client.post(
            reverse(
                "fundboard:add_virtual_watchlist",
                args=[portfolio.pk, self.stock.pk],
            ),
            {"allow_fractional": ""},
        )
        entry = portfolio.watchlist_entries.get(instrument=self.stock)
        self.assertEqual(add_watch.status_code, 302)
        self.assertFalse(entry.allow_fractional)

        limit_response = self.client.post(
            reverse("fundboard:place_virtual_order", args=[portfolio.pk]),
            {
                "instrument": self.stock.pk,
                "side": VirtualOrder.Side.BUY,
                "order_type": VirtualOrder.OrderType.LIMIT,
                "quantity": "1",
                "limit_price": "90",
            },
        )
        open_order = portfolio.virtual_orders.get()
        self.assertEqual(limit_response.status_code, 302)
        self.assertEqual(open_order.status, VirtualOrder.Status.OPEN)

        cancel_response = self.client.post(
            reverse(
                "fundboard:cancel_virtual_order",
                args=[portfolio.pk, open_order.pk],
            )
        )
        open_order.refresh_from_db()
        self.assertEqual(cancel_response.status_code, 302)
        self.assertEqual(open_order.status, VirtualOrder.Status.CANCELLED)

        clone_response = self.client.post(
            reverse("fundboard:clone_virtual_portfolio", args=[portfolio.pk]),
            {"name": "Copie via interface"},
        )
        clone = VirtualPortfolio.objects.get(name="Copie via interface")
        self.assertRedirects(
            clone_response,
            reverse("fundboard:virtual_portfolio_detail", args=[clone.pk]),
        )

        reset_response = self.client.post(
            reverse("fundboard:reset_virtual_portfolio", args=[portfolio.pk])
        )
        self.assertEqual(reset_response.status_code, 302)
        self.assertEqual(portfolio.virtual_orders.count(), 0)
        self.assertEqual(portfolio.watchlist_entries.count(), 1)

    def test_search_only_returns_shared_or_owned_active_instruments(self):
        mine = Instrument.objects.create(
            owner=self.alice,
            name="Recherche Alice",
            instrument_type=Instrument.Type.ETF,
            ticker="ALICE",
            currency="EUR",
        )

        results = list(searchable_virtual_instruments(self.alice, "Recherche"))
        all_accessible = list(searchable_virtual_instruments(self.alice))

        self.assertEqual(results, [mine])
        self.assertIn(self.stock, all_accessible)
        self.assertNotIn(self.bob_instrument, all_accessible)
