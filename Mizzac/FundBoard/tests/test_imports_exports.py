import json
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from FundBoard.exports.service import excel_export_bytes, json_export_bytes
from FundBoard.imports.parser import parse_json_document, parse_xlsx_document
from FundBoard.imports.service import (
    DuplicateImportError,
    RollbackRefusedError,
    commit_import,
    preview_import,
    rollback_import,
)
from FundBoard.imports.templates import excel_template_bytes, json_template_bytes
from FundBoard.models import Account, FinancialAuditEvent, ImportBatch, Instrument


def json_bytes(**sections):
    return json.dumps({"schema_version": "1.0", **sections}).encode()


class ImportServiceTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="import-alice")

    def test_partial_import_keeps_valid_row_and_rejects_invalid_row(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "valid-account",
                    "name": "Compte valide",
                    "category": "CURRENT",
                    "currency": "EUR",
                },
                {
                    "manual_reference": "invalid-account",
                    "name": "Compte invalide",
                    "category": "NOT_A_CATEGORY",
                    "currency": "EUR",
                },
            ]
        )
        document = parse_json_document(content)

        report = commit_import(
            self.user,
            document,
            file_name="partial.json",
            file_format="JSON",
            content=content,
        )

        self.assertEqual(report.created_rows, 1)
        self.assertEqual(report.invalid_rows, 1)
        self.assertEqual(report.batch.status, ImportBatch.Status.PARTIAL)
        self.assertTrue(Account.objects.filter(name="Compte valide").exists())
        self.assertFalse(Account.objects.filter(name="Compte invalide").exists())
        self.assertTrue(
            report.batch.issues.filter(sheet="Accounts", row_number=3, column="category").exists()
        )
        self.assertTrue(
            FinancialAuditEvent.objects.filter(
                user=self.user,
                event_type=FinancialAuditEvent.Type.IMPORT,
                object_pk=report.batch.pk,
            ).exists()
        )

    def test_parser_error_row_is_not_executed(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "unknown-column",
                    "name": "À ignorer",
                    "category": "CURRENT",
                    "currency": "EUR",
                    "secret": "ne doit pas entrer",
                }
            ]
        )
        document = parse_json_document(content)

        report = commit_import(
            self.user,
            document,
            file_name="unknown.json",
            file_format="JSON",
            content=content,
        )

        self.assertEqual(report.created_rows, 0)
        self.assertEqual(report.invalid_rows, 1)
        self.assertEqual(report.batch.status, ImportBatch.Status.FAILED)
        self.assertFalse(Account.objects.filter(user=self.user).exists())

    def test_dry_run_rolls_back_every_write(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "preview-account",
                    "name": "Prévisualisation",
                    "category": "SAVINGS",
                    "currency": "EUR",
                }
            ]
        )

        report = preview_import(self.user, parse_json_document(content))

        self.assertEqual(report.created_rows, 1)
        self.assertFalse(Account.objects.filter(user=self.user).exists())
        self.assertFalse(ImportBatch.objects.filter(user=self.user).exists())

    def test_duplicate_file_is_blocked_then_allowed_after_rollback(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "rollback-account",
                    "name": "À annuler",
                    "category": "CURRENT",
                    "currency": "EUR",
                }
            ]
        )
        document = parse_json_document(content)
        first = commit_import(
            self.user,
            document,
            file_name="same.json",
            file_format="JSON",
            content=content,
        )

        with self.assertRaises(DuplicateImportError):
            commit_import(
                self.user,
                document,
                file_name="same-again.json",
                file_format="JSON",
                content=content,
            )

        rollback_import(first.batch)
        self.assertFalse(Account.objects.filter(manual_reference="rollback-account").exists())
        self.assertTrue(
            FinancialAuditEvent.objects.filter(
                user=self.user,
                event_type=FinancialAuditEvent.Type.IMPORT_ROLLBACK,
                object_pk=first.batch.pk,
            ).exists()
        )
        second = commit_import(
            self.user,
            document,
            file_name="same-after-rollback.json",
            file_format="JSON",
            content=content,
        )
        self.assertEqual(second.created_rows, 1)

    def test_rollback_refuses_an_object_changed_after_import(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "changed-account",
                    "name": "Avant",
                    "category": "CURRENT",
                    "currency": "EUR",
                }
            ]
        )
        report = commit_import(
            self.user,
            parse_json_document(content),
            file_name="changed.json",
            file_format="JSON",
            content=content,
        )
        account = Account.objects.get(manual_reference="changed-account")
        account.name = "Après une modification manuelle"
        account.save(update_fields=["name"])

        with self.assertRaises(RollbackRefusedError):
            rollback_import(report.batch)

        report.batch.refresh_from_db()
        self.assertEqual(report.batch.status, ImportBatch.Status.COMPLETED)


class ImportTemplateTests(TestCase):
    def test_json_and_excel_templates_are_versioned_and_parseable(self):
        json_document = parse_json_document(json_template_bytes())
        excel_content = excel_template_bytes()
        excel_document = parse_xlsx_document(excel_content)
        workbook = load_workbook(BytesIO(excel_content))

        self.assertEqual(json_document.schema_version, "1.0")
        self.assertEqual(excel_document.schema_version, "1.0")
        self.assertEqual(
            workbook.sheetnames,
            [
                "README",
                "Lists",
                "Accounts",
                "Instruments",
                "Positions",
                "Transactions",
                "RealEstate",
                "Loans",
            ],
        )
        self.assertEqual(workbook["Lists"].sheet_state, "hidden")
        self.assertGreater(len(workbook["Accounts"].data_validations.dataValidation), 0)

    def test_excel_formula_row_is_rejected(self):
        workbook = load_workbook(BytesIO(excel_template_bytes()))
        workbook["Accounts"]["B2"] = "=1+1"
        output = BytesIO()
        workbook.save(output)

        document = parse_xlsx_document(output.getvalue())

        self.assertFalse(
            any(row.sheet == "Accounts" and row.row_number == 2 for row in document.rows)
        )
        self.assertTrue(
            any(issue.code == "formula_not_allowed" for issue in document.issues)
        )


class ExportAndImportViewsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.alice = get_user_model().objects.create_user(username="export-alice")
        cls.bob = get_user_model().objects.create_user(username="export-bob")
        cls.account = Account.objects.create(
            user=cls.alice,
            manual_reference="alice-account",
            name="Compte Alice",
            category=Account.Type.CURRENT,
            currency="EUR",
            metadata={"token": "never-export"},
        )
        Account.objects.create(
            user=cls.bob,
            manual_reference="bob-account",
            name="Compte Bob secret",
            category=Account.Type.CURRENT,
            currency="EUR",
        )
        Instrument.objects.create(
            owner=cls.alice,
            manual_reference="alice-instrument",
            name="Instrument Alice",
            instrument_type=Instrument.Type.ETF,
            currency="EUR",
            provider_identifiers={"provider": "secret-payload"},
        )

    def setUp(self):
        self.client.force_login(self.alice)

    def test_exports_are_owner_scoped_and_exclude_sensitive_payloads(self):
        payload = json.loads(json_export_bytes(self.alice))
        serialized = json.dumps(payload)

        self.assertIn("Compte Alice", serialized)
        self.assertNotIn("Compte Bob secret", serialized)
        self.assertNotIn("never-export", serialized)
        self.assertNotIn("secret-payload", serialized)
        self.assertNotIn("metadata", serialized)
        self.assertNotIn("provider_identifiers", serialized)
        self.assertGreater(len(excel_export_bytes(self.alice, "accounts")), 1000)

        response = self.client.get(
            reverse("fundboard:export", args=["json"]),
            {"resource": "accounts"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(
            FinancialAuditEvent.objects.filter(
                user=self.alice,
                event_type=FinancialAuditEvent.Type.EXPORT,
                details={"format": "json", "resource": "accounts"},
            ).exists()
        )

    def test_import_history_is_private_and_errors_are_downloadable(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "bad",
                    "name": "Invalide",
                    "category": "INVALID",
                    "currency": "EUR",
                }
            ]
        )
        report = commit_import(
            self.alice,
            parse_json_document(content),
            file_name="errors.json",
            file_format="JSON",
            content=content,
        )

        own = self.client.get(reverse("fundboard:import_batch", args=[report.batch.pk]))
        csv_response = self.client.get(
            reverse("fundboard:import_issues_csv", args=[report.batch.pk])
        )
        self.client.force_login(self.bob)
        other = self.client.get(reverse("fundboard:import_batch", args=[report.batch.pk]))

        self.assertEqual(own.status_code, 200)
        self.assertContains(own, "INVALID")
        self.assertEqual(csv_response.status_code, 200)
        self.assertTrue(csv_response.content.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(other.status_code, 404)

    def test_upload_preview_does_not_write(self):
        content = json_bytes(
            accounts=[
                {
                    "manual_reference": "web-preview",
                    "name": "Web preview",
                    "category": "CURRENT",
                    "currency": "EUR",
                }
            ]
        )
        upload = SimpleUploadedFile("preview.json", content, content_type="application/json")

        response = self.client.post(
            reverse("fundboard:import_export"),
            {"action": "preview", "file": upload},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Aucune écriture")
        self.assertFalse(Account.objects.filter(manual_reference="web-preview").exists())
