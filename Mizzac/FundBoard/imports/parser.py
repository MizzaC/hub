import json
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .schema import MAX_IMPORT_BYTES, SCHEMA_VERSION, SHEETS

MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_MEMBERS = 1000


class ImportDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedRow:
    sheet: str
    row_number: int
    data: dict


@dataclass(frozen=True)
class RowIssue:
    sheet: str
    row_number: int
    column: str
    code: str
    message: str
    value_preview: str = ""


@dataclass(frozen=True)
class ParsedDocument:
    schema_version: str
    rows: tuple[ParsedRow, ...]
    issues: tuple[RowIssue, ...]


def _safe_preview(column, value):
    if column == "address":
        return "[masqué]"
    return str(value)[:200] if value is not None else ""


def _validate_row(sheet, row_number, data):
    config = SHEETS[sheet]
    issues = []
    for column, value in data.items():
        if isinstance(value, str) and value.lstrip().startswith("="):
            issues.append(
                RowIssue(
                    sheet,
                    row_number,
                    column,
                    "formula_not_allowed",
                    "Les formules ne sont pas autorisées dans les imports.",
                )
            )
    unknown = set(data) - set(config["columns"])
    for column in sorted(unknown):
        issues.append(
            RowIssue(
                sheet,
                row_number,
                column,
                "unknown_column",
                "Colonne non reconnue par le schéma 1.0.",
                _safe_preview(column, data[column]),
            )
        )
    for column in sorted(config["required"]):
        if data.get(column) in (None, ""):
            issues.append(
                RowIssue(sheet, row_number, column, "required", "Valeur obligatoire manquante.")
            )
    return issues


def _validate_xlsx_archive(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise ImportDocumentError("Le classeur Excel contient trop de fichiers internes.")
            if sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ImportDocumentError("Le classeur Excel décompressé dépasse 50 Mio.")
            lowered_names = {member.filename.lower() for member in members}
            if any("vbaproject.bin" in name for name in lowered_names):
                raise ImportDocumentError("Les macros Excel ne sont pas autorisées.")
            if any("externallinks/" in name for name in lowered_names):
                raise ImportDocumentError("Les liens Excel externes ne sont pas autorisés.")
    except BadZipFile as exc:
        raise ImportDocumentError("Le classeur Excel est invalide.") from exc


def parse_json_document(content):
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportDocumentError("Le fichier dépasse 5 Mio.")
    try:
        payload = json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ImportDocumentError("Le document JSON est invalide ou n'est pas encodé en UTF-8.") from exc
    if not isinstance(payload, dict):
        raise ImportDocumentError("La racine JSON doit être un objet.")
    if payload.get("schema_version") != SCHEMA_VERSION:
        raise ImportDocumentError(f"schema_version doit valoir {SCHEMA_VERSION}.")

    allowed_keys = {"schema_version"} | {config["json_key"] for config in SHEETS.values()}
    unknown_keys = set(payload) - allowed_keys
    if unknown_keys:
        raise ImportDocumentError(f"Sections JSON inconnues : {', '.join(sorted(unknown_keys))}.")

    rows = []
    issues = []
    for sheet, config in SHEETS.items():
        values = payload.get(config["json_key"], [])
        if not isinstance(values, list):
            issues.append(RowIssue(sheet, 0, "", "invalid_section", "La section doit être une liste."))
            continue
        for index, value in enumerate(values, start=2):
            if not isinstance(value, dict):
                issues.append(RowIssue(sheet, index, "", "invalid_row", "La ligne doit être un objet JSON."))
                continue
            row_issues = _validate_row(sheet, index, value)
            issues.extend(row_issues)
            if not row_issues:
                rows.append(ParsedRow(sheet, index, value))
    return ParsedDocument(SCHEMA_VERSION, tuple(rows), tuple(issues))


def parse_xlsx_document(content):
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportDocumentError("Le fichier dépasse 5 Mio.")
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ImportDocumentError("Le classeur Excel est invalide.") from exc

    version = workbook["README"]["B1"].value if "README" in workbook.sheetnames else None
    if str(version) != SCHEMA_VERSION:
        raise ImportDocumentError(f"La cellule README!B1 doit contenir la version {SCHEMA_VERSION}.")

    rows = []
    issues = []
    for sheet, config in SHEETS.items():
        if sheet not in workbook.sheetnames:
            issues.append(RowIssue(sheet, 0, "", "missing_sheet", "Feuille obligatoire absente."))
            continue
        worksheet = workbook[sheet]
        iterator = worksheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            issues.append(RowIssue(sheet, 1, "", "missing_header", "Ligne d'en-tête absente."))
            continue
        normalized_headers = [str(value).strip() if value is not None else "" for value in headers]
        unknown = set(normalized_headers) - set(config["columns"]) - {""}
        if unknown:
            issues.append(
                RowIssue(
                    sheet,
                    1,
                    ", ".join(sorted(unknown)),
                    "unknown_column",
                    "Le classeur contient une colonne inconnue.",
                )
            )
            continue
        header_positions = [(index, header) for index, header in enumerate(normalized_headers) if header]
        for row_number, values in enumerate(iterator, start=2):
            data = {
                header: values[index] if index < len(values) else None
                for index, header in header_positions
            }
            if all(value in (None, "") for value in data.values()):
                continue
            row_issues = _validate_row(sheet, row_number, data)
            issues.extend(row_issues)
            if not row_issues:
                rows.append(ParsedRow(sheet, row_number, data))
    workbook.close()
    return ParsedDocument(SCHEMA_VERSION, tuple(rows), tuple(issues))


def parse_document(content, file_name):
    suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    if suffix == "json":
        return parse_json_document(content), "JSON"
    if suffix == "xlsx":
        return parse_xlsx_document(content), "XLSX"
    raise ImportDocumentError("Formats acceptés : .json et .xlsx.")
