import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .schema import CHOICE_LISTS, SCHEMA_VERSION, SHEETS, VALIDATION_COLUMNS

HEADER_FILL = PatternFill("solid", fgColor="206BC4")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INVALID_FILL = PatternFill("solid", fgColor="FDE2E2")


def json_template_bytes():
    payload = {"schema_version": SCHEMA_VERSION}
    for config in SHEETS.values():
        payload[config["json_key"]] = [dict(zip(config["columns"], config["example"]))]
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def excel_template_bytes():
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["schema_version", SCHEMA_VERSION])
    readme.append(["Mode d'import", "Succès partiel : chaque ligne invalide est ignorée et rapportée."])
    readme.append(["Références", "Utilisez des références personnelles stables entre les feuilles."])
    readme.append(["Sécurité", "N'ajoutez jamais de mot de passe, token, cookie, PIN, seed ou clé privée."])
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 90
    readme["A1"].font = Font(bold=True)

    lists = workbook.create_sheet("Lists")
    list_ranges = {}
    for column_index, (name, values) in enumerate(CHOICE_LISTS.items(), start=1):
        lists.cell(1, column_index, name)
        lists.cell(1, column_index).font = Font(bold=True)
        for row_index, value in enumerate(values, start=2):
            lists.cell(row_index, column_index, value)
        letter = get_column_letter(column_index)
        workbook.defined_names.add(
            DefinedName(name, attr_text=f"'Lists'!${letter}$2:${letter}${len(values) + 1}")
        )
        list_ranges[name] = name
    lists.sheet_state = "hidden"

    for sheet, config in SHEETS.items():
        worksheet = workbook.create_sheet(sheet)
        worksheet.append(config["columns"])
        worksheet.append(config["example"])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(config['columns']))}2"
        for column_index, column in enumerate(config["columns"], start=1):
            cell = worksheet.cell(1, column_index)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(len(column) + 3, 14),
                28,
            )
            choice_name = VALIDATION_COLUMNS.get((sheet, column))
            if choice_name:
                validation = DataValidation(
                    type="list",
                    formula1=f"={list_ranges[choice_name]}",
                    allow_blank=column not in config["required"],
                )
                validation.error = "Valeur absente de la liste autorisée."
                validation.errorTitle = "Valeur invalide"
                validation.prompt = "Sélectionnez une valeur de la liste."
                validation.promptTitle = column
                validation.showErrorMessage = True
                validation.showInputMessage = True
                worksheet.add_data_validation(validation)
                letter = get_column_letter(column_index)
                validation.add(f"{letter}2:{letter}10000")
        for required in config["required"]:
            column_index = config["columns"].index(required) + 1
            letter = get_column_letter(column_index)
            worksheet.conditional_formatting.add(
                f"{letter}2:{letter}10000",
                FormulaRule(formula=[f'LEN(TRIM({letter}2))=0'], fill=INVALID_FILL),
            )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
