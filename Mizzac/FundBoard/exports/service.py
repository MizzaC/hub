import csv
import json
from io import BytesIO, StringIO

from django.db.models import Q
from openpyxl import load_workbook

from FundBoard.imports.schema import SCHEMA_VERSION, SHEETS
from FundBoard.imports.templates import excel_template_bytes
from FundBoard.models import Account, Instrument, Loan, Position, RealEstate, Transaction

RESOURCE_NAMES = {
    "accounts": "Accounts",
    "instruments": "Instruments",
    "positions": "Positions",
    "transactions": "Transactions",
    "real_estate": "RealEstate",
    "loans": "Loans",
}

NUMERIC_COLUMNS = {
    "balance",
    "ownership_share",
    "quantity",
    "average_unit_cost",
    "cost_basis",
    "current_unit_price",
    "current_value",
    "unit_price",
    "gross_amount",
    "fees",
    "taxes",
    "net_amount",
    "surface_sqm",
    "purchase_price",
    "purchase_costs",
    "renovation_costs",
    "estimated_value",
    "monthly_rent",
    "monthly_charges",
    "annual_property_tax",
    "annual_insurance",
    "annual_other_costs",
    "vacancy_rate",
    "original_principal",
    "outstanding_principal",
    "nominal_rate",
    "apr",
    "duration_months",
    "payment_amount",
    "insurance_amount",
    "insurance_rate",
    "initial_fees",
    "deferred_months",
}


def _value(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value) if value.__class__.__name__ == "Decimal" else value


def _account_reference(account):
    return account.manual_reference or f"account-{account.pk}"


def _instrument_reference(instrument):
    return instrument.manual_reference or f"instrument-{instrument.pk}"


def _loan_reference(loan):
    return loan.manual_reference or f"loan-{loan.pk}"


def _excel_safe(value):
    if isinstance(value, str) and value.lstrip().startswith("="):
        return f"'{value}"
    return value


def _csv_safe(column, value):
    if (
        column not in NUMERIC_COLUMNS
        and isinstance(value, str)
        and value.lstrip().startswith(("=", "+", "-", "@"))
    ):
        return f"'{value}"
    return value


def export_rows(user):
    accounts = list(Account.objects.filter(user=user).order_by("pk"))
    instruments = list(
        Instrument.objects.filter(Q(owner=user) | Q(positions__account__user=user))
        .distinct()
        .order_by("pk")
    )
    positions = list(
        Position.objects.filter(account__user=user)
        .select_related("account", "instrument")
        .order_by("pk")
    )
    transactions = list(
        Transaction.objects.filter(user=user)
        .select_related("account", "instrument")
        .order_by("pk")
    )
    loans = list(Loan.objects.filter(user=user).select_related("account").order_by("pk"))
    real_estate = list(
        RealEstate.objects.filter(user=user).select_related("linked_loan").order_by("pk")
    )
    return {
        "Accounts": [
            {
                "manual_reference": _account_reference(item),
                "name": item.name,
                "category": item.category,
                "subtype": item.subtype,
                "balance": _value(item.balance),
                "currency": item.currency,
                "ownership_share": _value(item.ownership_share),
                "iban_masked": item.iban_masked,
                "status": item.status,
            }
            for item in accounts
        ],
        "Instruments": [
            {
                "manual_reference": _instrument_reference(item),
                "name": item.name,
                "instrument_type": item.instrument_type,
                "ticker": item.ticker,
                "isin": item.isin,
                "market_mic": item.market_mic,
                "currency": item.currency,
                "country_code": item.country_code,
                "sector": item.sector,
                "blockchain": item.blockchain,
                "contract_address": item.contract_address,
                "status": item.status,
            }
            for item in instruments
        ],
        "Positions": [
            {
                "account_reference": _account_reference(item.account),
                "instrument_reference": _instrument_reference(item.instrument),
                "quantity": _value(item.quantity),
                "average_unit_cost": _value(item.average_unit_cost),
                "cost_basis": _value(item.cost_basis),
                "current_unit_price": _value(item.current_unit_price),
                "current_value": _value(item.current_value),
                "value_currency": item.value_currency,
                "valued_at": _value(item.valued_at),
                "status": item.status,
            }
            for item in positions
        ],
        "Transactions": [
            {
                "account_reference": _account_reference(item.account),
                "instrument_reference": (
                    _instrument_reference(item.instrument) if item.instrument else ""
                ),
                "transaction_type": item.transaction_type,
                "subtype": item.subtype,
                "quantity": _value(item.quantity),
                "unit_price": _value(item.unit_price),
                "gross_amount": _value(item.gross_amount),
                "fees": _value(item.fees),
                "taxes": _value(item.taxes),
                "net_amount": _value(item.net_amount),
                "currency": item.currency,
                "executed_at": _value(item.executed_at),
                "value_date": _value(item.value_date),
                "label": item.label,
                "status": item.status,
                "idempotency_key": item.idempotency_key,
            }
            for item in transactions
        ],
        "RealEstate": [
            {
                "manual_reference": item.manual_reference,
                "name": item.name,
                "property_type": item.property_type,
                "address": item.address,
                "surface_sqm": _value(item.surface_sqm),
                "ownership_share": _value(item.ownership_share),
                "currency": item.currency,
                "purchase_price": _value(item.purchase_price),
                "purchase_costs": _value(item.purchase_costs),
                "renovation_costs": _value(item.renovation_costs),
                "estimated_value": _value(item.estimated_value),
                "valuation_date": _value(item.valuation_date),
                "valuation_source": item.valuation_source,
                "monthly_rent": _value(item.monthly_rent),
                "monthly_charges": _value(item.monthly_charges),
                "annual_property_tax": _value(item.annual_property_tax),
                "annual_insurance": _value(item.annual_insurance),
                "annual_other_costs": _value(item.annual_other_costs),
                "vacancy_rate": _value(item.vacancy_rate),
                "loan_reference": _loan_reference(item.linked_loan) if item.linked_loan else "",
            }
            for item in real_estate
        ],
        "Loans": [
            {
                "manual_reference": item.manual_reference,
                "name": item.name,
                "loan_type": item.loan_type,
                "account_reference": _account_reference(item.account) if item.account else "",
                "currency": item.currency,
                "original_principal": _value(item.original_principal),
                "outstanding_principal": _value(item.outstanding_principal),
                "balance_date": _value(item.balance_date),
                "nominal_rate": _value(item.nominal_rate),
                "apr": _value(item.apr),
                "duration_months": item.duration_months,
                "start_date": _value(item.start_date),
                "maturity_date": _value(item.maturity_date),
                "payment_amount": _value(item.payment_amount),
                "insurance_amount": _value(item.insurance_amount),
                "insurance_rate": _value(item.insurance_rate),
                "initial_fees": _value(item.initial_fees),
                "deferred_months": item.deferred_months,
            }
            for item in loans
        ],
    }


def json_export_bytes(user, resource="all"):
    rows = export_rows(user)
    payload = {"schema_version": SCHEMA_VERSION}
    for sheet, config in SHEETS.items():
        if resource == "all" or RESOURCE_NAMES.get(resource) == sheet:
            payload[config["json_key"]] = rows[sheet]
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def excel_export_bytes(user, resource="all"):
    rows = export_rows(user)
    workbook = load_workbook(BytesIO(excel_template_bytes()))
    for sheet, config in SHEETS.items():
        worksheet = workbook[sheet]
        worksheet.delete_rows(2, worksheet.max_row - 1)
        if resource != "all" and RESOURCE_NAMES.get(resource) != sheet:
            continue
        for item in rows[sheet]:
            worksheet.append([_excel_safe(item.get(column)) for column in config["columns"]])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def csv_export_bytes(user, resource):
    sheet = RESOURCE_NAMES.get(resource)
    if not sheet:
        raise ValueError("Une ressource précise est obligatoire pour l'export CSV.")
    config = SHEETS[sheet]
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=config["columns"])
    writer.writeheader()
    writer.writerows(
        {
            column: _csv_safe(column, item.get(column))
            for column in config["columns"]
        }
        for item in export_rows(user)[sheet]
    )
    return output.getvalue().encode("utf-8-sig")
